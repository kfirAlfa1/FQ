import API
import ReadData
from openpyxl import load_workbook
from openpyxl import Workbook
import tkinter as tk
from tkinter import filedialog
import Features
import utils
import TokenAndCats

token_and_cats = TokenAndCats.TokenAndCats()
sheets_array = []

def load_and_convert_excel_to_values_only(num_sheets):
    global sheets_array

    file_path = filedialog.askopenfilename(
        title="Select an Excel file",
        filetypes=[("Excel files", "*.xlsx;*.xls")]
    )

    if file_path:
        wb = load_workbook(file_path, data_only=True)  # Open Excel file (values only)
        new_wb = Workbook()

        # Copy each worksheet into the new workbook
        for idx, sheet_name in enumerate(wb.sheetnames):
            original_ws = wb[sheet_name]
            if idx == 0:
                new_ws = new_wb.active
            else:
                new_ws = new_wb.create_sheet(title=sheet_name)

            new_ws.title = sheet_name

            for row in original_ws.iter_rows(values_only=True):
                new_ws.append(row)

        sheets_array = new_wb.worksheets[:num_sheets]

def create_systems(company_id, sheets):
    if company_id == None:
        ReadData.company_name=ReadData.get_company_name(sheets[0])
        API.create_company()
    else:
        API.company_id=company_id

    for sheet in sheets:
        create_system(sheet)

def create_system(sheet):
    ReadData.read_all_data(sheet)
    API.create_app()
    token_and_cats.get_cats(API.app_id)
    API.create_roles(ReadData.roles, API.app_id)
    API.create_shifts(ReadData.shifts, API.app_id)
    API.create_requests_limits(ReadData.requests_limits, API.app_id)
    API.create_recives_limits(ReadData.receives_limits, API.app_id)
    API.create_clock(API.app_id)
    API.create_features(ReadData.features, API.app_id)
    if len(ReadData.cutting_hours) != 0:
        Features.create_cutting_hours_feature(API.app_id, ReadData.cutting_hours)
    if ReadData.shabat is not None and len(ReadData.shabat) != 0:
        API.create_shabat(API.app_id)
    if ReadData.holidays == "כן":
        ReadData.holidays_area = utils.choose_holiday_area_code(ReadData.holidays_area)
        API.create_automatic_holiday(API.app_id)
    if ReadData.automatic_holidays_special_cost == "כן":
        API.create_report_attribute(API.app_id)
        Features.create_special_holiday(API.app_id)
        if ReadData.language != "he":
            API.create_report_coulmn(API.app_id, ["1", "2"], 'special_holiday_100', "Speciel 200%")
        else:
            API.create_report_coulmn(API.app_id, ["1", "2"], 'special_holiday_100', "מיוחד 200%")
    API.create_employee_configurations(ReadData.employee_configurations, API.app_id)
    API.create_employees(ReadData.owner_details, API.app_id)
    
    print("\n\nSystem created successfully\n\n")

def create_attendance_systems(company_id, sheets):
    if company_id == None:
        ReadData.company_name=ReadData.get_company_name(sheets[0])
        API.create_company()
    else:
        API.company_id=company_id

    for sheet in sheets:
        create_attendance_system(sheet)

def create_attendance_system(sheet):
    ReadData.read_all_data(sheet)
    API.create_attendance_app()
    token_and_cats.get_cats(API.app_id)
    API.create_roles(ReadData.roles, API.app_id)
    unplanned_shift_flag = True
    API.create_unplanned_shift(API.app_id, unplanned_shift_flag)
    API.create_clock(API.app_id)
    API.create_features(ReadData.features, API.app_id)
    if len(ReadData.cutting_hours) != 0:
        Features.create_cutting_hours_feature(API.app_id, ReadData.cutting_hours)
    if ReadData.shabat is not None and len(ReadData.shabat) != 0:
        API.create_shabat(API.app_id)
    if ReadData.automatic_holidays_special_cost == "כן":
        API.create_report_attribute(API.app_id)
        Features.create_special_holiday(API.app_id)
        if ReadData.language != "he":
            API.create_report_coulmn(API.app_id, ["1", "2"], 'special_holiday_100', "Speciel 200%")
        else:
            API.create_report_coulmn(API.app_id, ["1", "2"], 'special_holiday_100', "מיוחד 200%")
    if ReadData.holidays == "כן":
        ReadData.holidays_area = utils.choose_holiday_area_code(ReadData.holidays_area)
        API.create_automatic_holiday(API.app_id)
    Features.delete_auto_populate_rota(API.app_id)
    API.create_employee_configurations(ReadData.employee_configurations, API.app_id)
    API.create_employees(ReadData.owner_details, API.app_id)

    print("\n\nAttendance system created successfully\n\n")

def execute_system():
    company_id_text = company_id_entry.get().strip()
    company_id = int(company_id_text) if company_id_text else None

    if not sheets_array:
        print("Please load an Excel file first.")
        return

    selection = system_choice.get()
    if selection == 1:
        create_systems(company_id, sheets_array)
    elif selection == 2:
        create_attendance_systems(company_id, sheets_array)
    else:
        print("Invalid selection.")

def on_load_button_click():
    try:
        num_sheets = int(num_sheets_entry.get())
    except ValueError:
        print("Please enter a valid integer for number of sheets.")
        return

    load_and_convert_excel_to_values_only(num_sheets)



# ----- GUI Setup -----
#######################

root = tk.Tk()
root.title("FQ system")
root.geometry("400x300")

# Input for number of sheets
tk.Label(root, text="Number of sheets to load:").pack()
num_sheets_entry = tk.Entry(root)
num_sheets_entry.pack(pady=5)

# Input for company ID
tk.Label(root, text="Company ID:").pack()
company_id_entry = tk.Entry(root)
company_id_entry.pack(pady=5)

# Button to load the Excel file.
# Updated to use the on_load_button_click wrapper which processes the input value.
load_button = tk.Button(root, text="Load Excel File", command=on_load_button_click)
load_button.pack(pady=10)

# Two radio buttons for the two options.
system_choice = tk.IntVar()
system_choice.set(1)  # Default selection

tk.Radiobutton(root, text="Option 1: Create Systems", variable=system_choice, value=1).pack(anchor="w", padx=20)
tk.Radiobutton(root, text="Option 2: Create Attendance System", variable=system_choice, value=2).pack(anchor="w", padx=20)

# Execute button.
execute_button = tk.Button(root, text="Execute", command=lambda: execute_system(), width=20, height=2, font=("Helvetica", 14))
execute_button.pack(pady=10)

root.mainloop()
###################################
###################################

#For Motti
# Support gruop in  organize_groups() 
